"""実利用Excelを専用DBで検証する。通常保存先・外部APIは使用しない。"""
from __future__ import annotations

import argparse
from collections import Counter
from contextlib import closing
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import sqlite3
import sys
import unicodedata
import csv

# このファイルはプロジェクトのscriptsへ配置する。
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from openpyxl import load_workbook
from book_manager.database.initializer import initialize_database
from book_manager.models.excel_import import ExcelImportOptions
from book_manager.repositories.book_repository import BookRepository
from book_manager.services.backup_service import BackupService
from book_manager.services.book_service import BookService
from book_manager.services.excel_import_preflight_service import ExcelImportPreflightService
from book_manager.services.excel_import_service import ExcelImportService

HEADERS = ('書籍名', '著者名', '出版社', '出版形態', '開始月', '読了月', 'あらすじ')
MAPPING = (('title', 1), ('author', 2), ('publisher', 3), ('publication_format', 4),
           ('started_on', 5), ('completed_on', 6), ('description', 7))


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def clean(value):
    return value.strip() or None if isinstance(value, str) else value


def expected_month(value):
    """取込側の日付変換を使わず、このExcelの年月表記を独立に検証する。"""
    if value is None or not str(value).strip():
        return None
    if not isinstance(value, str):
        raise ValueError('This verifier expects month strings in columns E/F.')
    text = re.sub(r'\s+', '', unicodedata.normalize('NFKC', value))
    match = re.fullmatch(r'(昭和|平成|令和)(元|\d+)年(\d{1,2})月', text)
    if match:
        era, year, month = match.groups()
        year = {'昭和': 1925, '平成': 1988, '令和': 2018}[era] + (1 if year == '元' else int(year))
    else:
        match = re.fullmatch(r'(\d{4})(?:年|[-/.])(\d{1,2})月?', text)
        if not match:
            raise ValueError('Unsupported month value: ' + repr(value))
        year, month = match.groups()
    return date(int(year), int(month), 1).strftime('%Y-%m')


def read_source(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Exactly one worksheet is required.')
        sheet = workbook.worksheets[0]
        if tuple(c.value for c in next(sheet.iter_rows(max_row=1))) != HEADERS:
            raise ValueError('Unexpected headers. Do not guess column mapping.')
        records, skipped = [], []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            values = tuple(c.value for c in cells)
            if all(v is None or isinstance(v, str) and not v.strip() for v in values):
                skipped.append(row_number)
                continue
            if any(c.data_type == 'f' for c in cells):
                raise ValueError('Formula cells require separate validation.')
            title, author, publisher, form, started, completed, description = values
            if not isinstance(title, str) or not title.strip():
                raise ValueError('Missing title at row ' + str(row_number))
            start, finish = expected_month(started), expected_month(completed)
            records.append((row_number, dict(
                title=clean(title), author=clean(author), publisher=clean(publisher),
                publication_format=clean(form), description=clean(description),
                completion_month=finish, reading_status='completed' if finish else 'reading' if start else None,
                started_on=start, completed_on=finish,
            )))
        return sheet.title, records, skipped
    finally:
        workbook.close()


def readonly(path):
    return sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)


def verify(source, db):
    checks, failures = [], []
    with closing(readonly(db)) as connection:
        connection.row_factory = sqlite3.Row
        books = connection.execute('SELECT * FROM book_items ORDER BY id').fetchall()
        integrity = [r[0] for r in connection.execute('PRAGMA integrity_check')]
        foreign_keys = connection.execute('PRAGMA foreign_key_check').fetchall()
        schema = connection.execute('PRAGMA user_version').fetchone()[0]
        histories = connection.execute('SELECT * FROM reading_histories ORDER BY id').fetchall()
        by_book = {}
        for h in histories:
            by_book.setdefault(h['book_item_id'], []).append(h)
        counts = {table: connection.execute('SELECT count(*) FROM ' + table).fetchone()[0]
                  for table in ('book_items', 'reading_histories', 'tags', 'book_item_tags')}
        if len(books) != len(source): failures.append('Book count mismatch')
        expected_histories = sum(bool(v['started_on'] or v['completed_on']) for _, v in source)
        if len(histories) != expected_histories: failures.append('History count mismatch')
        if integrity != ['ok'] or foreign_keys: failures.append('SQLite integrity or foreign-key failure')
        if counts['tags'] or counts['book_item_tags']: failures.append('Unexpected tags')
        for (excel_row, expected), book in zip(source, books):
            wrong = []
            for field in ('title', 'author', 'publisher', 'publication_format', 'description', 'completion_month', 'reading_status'):
                if book[field] != expected[field]: wrong.append(field)
            for field in ('isbn', 'reference_isbn', 'magazine_code', 'cover_image_url', 'cover_image_path', 'rating', 'notes'):
                if book[field] is not None: wrong.append(field)
            if book['is_favorite'] != 0: wrong.append('is_favorite')
            history = by_book.get(book['id'], [])
            has_history = bool(expected['started_on'] or expected['completed_on'])
            if len(history) != int(has_history):
                wrong.append('history_count')
            elif has_history:
                for field in ('started_on', 'completed_on'):
                    if history[0][field] != expected[field]: wrong.append(field)
                if history[0]['notes'] is not None: wrong.append('history_notes')
            checks.append((excel_row, book['id'], 'FAIL' if wrong else 'PASS', '|'.join(wrong)))
            if wrong: failures.append('Excel row ' + str(excel_row) + ': ' + ', '.join(wrong))
        return dict(counts=counts, schema_version=schema, integrity_check=integrity,
                    foreign_key_errors=len(foreign_keys), checked_rows=len(checks), failures=failures), checks


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--excel', required=True, type=Path)
    parser.add_argument('--work-dir', required=True, type=Path,
                        help='Dedicated test home. prepare/import require a NEW directory.')
    parser.add_argument('--mode', choices=('prepare', 'import', 'verify'), default='import')
    args = parser.parse_args()
    source_path, work = args.excel.resolve(), args.work_dir.resolve()
    if not source_path.is_file(): raise FileNotFoundError(source_path)
    if work == source_path.parent or work == Path(__file__).resolve().parents[1]:
        raise ValueError('Use a separate dedicated work directory.')
    source_hash = sha256(source_path)
    sheet_name, source, skipped = read_source(source_path)
    db = work / 'data' / 'books.db'
    if args.mode in ('prepare', 'import'):
        # Refuse existing directories: never initialize or import into a user DB.
        work.mkdir(parents=True, exist_ok=False)
        initialize_database(db)
    elif not db.is_file():
        raise FileNotFoundError(db)
    preparer = BookService(BookRepository(db))
    preflight = ExcelImportPreflightService(book_service=preparer)
    preview = preflight.preview_file(source_path, options=ExcelImportOptions(MAPPING))
    report = dict(time=datetime.now().astimezone().isoformat(timespec='seconds'),
                  mode=args.mode, source_sha256=source_hash, sheet=sheet_name,
                  source_rows=len(source), blank_rows=len(skipped),
                  preview_records=len(preview.records), preview_errors=len(preview.errors),
                  preview_warnings=dict(Counter(i.message for i in preview.warnings)),
                  statuses=dict(Counter(v['reading_status'] for _, v in source)),
                  formats=dict(Counter(v['publication_format'] for _, v in source)),
                  work_directory=str(work), database=str(db))
    report['errors'] = [dict(row=i.row_number, field=i.field, coordinate=i.coordinate, message=i.message) for i in preview.errors]
    if not preview.can_import:
        (work / 'verification_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        raise ValueError('Preflight failed. See verification_report.json')
    if len(preview.records) != len(source) or tuple(r.row_number for r in preview.records) != tuple(n for n, _ in source):
        raise ValueError('Preflight and independent source row counts differ.')
    if args.mode == 'import':
        backup = BackupService(database_path=db, backup_directory=work / 'backups')
        service = ExcelImportService(database_path=db, backup_service=backup, preflight_service=preflight)
        result = service.import_preview(preview)
        with closing(readonly(result.backup_path)) as connection:
            empty_backup = connection.execute('SELECT count(*) FROM book_items').fetchone()[0] == 0
            backup_integrity = connection.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
        if not empty_backup or not backup_integrity: raise ValueError('Pre-import backup validation failed.')
        report['backup'] = str(result.backup_path)
        report['backup_is_valid_empty_database'] = True
    if args.mode != 'prepare':
        outcome, checks = verify(source, db)
        report.update(outcome)
        with (work / 'row_checks.csv').open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow(('excel_row', 'book_id', 'result', 'mismatched_fields'))
            writer.writerows(checks)
        report['passed'] = not outcome['failures']
    report['source_unchanged'] = sha256(source_path) == source_hash
    if not report['source_unchanged']: raise ValueError('Excel file changed during validation.')
    (work / 'verification_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Mode:', args.mode)
    print('Work directory:', work)
    print('Source records:', len(source), '/ blank rows:', len(skipped))
    print('Preflight errors:', len(preview.errors), '/ warnings:', len(preview.warnings))
    if args.mode == 'prepare':
        print('READY: empty test DB prepared. Import through the GUI, then run --mode verify.')
    else:
        print('Checked records:', report['checked_rows'])
        print('Counts:', report['counts'])
        print('Result:', 'PASS' if report['passed'] else 'FAIL')
        for failure in report['failures'][:20]: print(failure)
    print('Source unchanged:', report['source_unchanged'])
    print('Report:', work / 'verification_report.json')
    return 0 if report.get('passed', True) else 1


if __name__ == '__main__':
    raise SystemExit(main())
