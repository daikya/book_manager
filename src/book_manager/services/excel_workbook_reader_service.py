"""1シートの.xlsxを読み取る。元ファイルの保存・DB操作・外部通信は行わない。"""
from contextlib import ExitStack
import hashlib
from io import BytesIO
from pathlib import Path
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from book_manager.models.excel_workbook import ExcelCell, ExcelRow, ExcelSheetData


class ExcelWorkbookReaderService:
    """セルの保存済み値と書式を、列対応前の情報として返す。"""
    MAX_FILE_SIZE = 32 * 1024 * 1024
    MAX_EXPANDED_SIZE = 128 * 1024 * 1024
    MAX_ARCHIVE_ENTRIES = 2048
    MAX_ROWS = 100_000
    MAX_COLUMNS = 256

    def read_sheet(self, source_path: Path | str, *, header_row: int = 1) -> ExcelSheetData:
        """1シートの指定行を見出しとして読み、以降の全行を保持する。

        数式の計算値不足・Excelエラーはセルのissueに記録する。
        未使用列の問題で取込全体を止めるかは、後続の列割り当てで判断する。
        """
        path = Path(source_path).resolve()
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Excelは.xlsx形式で指定してください。.xlsや.xlsmは対象外です。")
        if type(header_row) is not int or not 1 <= header_row <= self.MAX_ROWS:
            raise ValueError("見出し行は1以上の範囲内の整数で指定してください。")
        with path.open("rb") as source:
            content = source.read(self.MAX_FILE_SIZE + 1)
        if len(content) > self.MAX_FILE_SIZE:
            raise ValueError("Excelファイルが読み込み上限を超えています。")
        self._validate_archive(content)
        with ExitStack() as stack:
            formula_stream = stack.enter_context(BytesIO(content))
            value_stream = stack.enter_context(BytesIO(content))
            formulas = load_workbook(
                formula_stream, read_only=True, data_only=False, keep_links=False,
            )
            stack.callback(formulas.close)
            values = load_workbook(
                value_stream, read_only=True, data_only=True, keep_links=False,
            )
            stack.callback(values.close)
            if len(formulas.worksheets) != 1:
                raise ValueError("取り込み用Excelはワークシートを1枚だけにしてください。")
            sheet = formulas.worksheets[0]
            cached_sheet = values[sheet.title]
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            maximum_row = sheet.max_row or 1
            maximum_column = sheet.max_column or 1
            if maximum_row > self.MAX_ROWS or maximum_column > self.MAX_COLUMNS:
                raise ValueError("Excelの行数または列数が読み込み上限を超えています。")
            if header_row > maximum_row:
                raise ValueError("指定した見出し行がExcelに存在しません。")
            rows = []
            raw_rows = sheet.iter_rows(
                min_row=header_row, max_row=maximum_row, max_col=maximum_column,
            )
            cached_rows = cached_sheet.iter_rows(
                min_row=header_row, max_row=maximum_row, max_col=maximum_column,
            )
            for row_number, (raw_row, cached_row) in enumerate(
                zip(raw_rows, cached_rows, strict=True), start=header_row,
            ):
                cells = []
                for column, (raw_cell, cached_cell) in enumerate(
                    zip(raw_row, cached_row, strict=True), start=1,
                ):
                    is_formula = raw_cell.data_type == "f"
                    cell_value = cached_cell.value if is_formula else raw_cell.value
                    issue = None
                    if is_formula and cell_value is None:
                        issue = "数式の保存済み計算値がありません。Excelで再計算して保存してください。"
                    elif raw_cell.data_type == "e" or cached_cell.data_type == "e":
                        issue = "Excelセルがエラー値です。元のExcelで修正してください。"
                    cells.append(ExcelCell(
                        coordinate=f"{get_column_letter(column)}{row_number}",
                        value=cell_value, number_format=raw_cell.number_format or "General",
                        is_formula=is_formula, issue=issue,
                    ))
                rows.append(ExcelRow(row_number=row_number, cells=tuple(cells)))
            header = rows[0]
            if header.is_empty:
                raise ValueError("見出し行が空です。正しい見出し行を指定してください。")
            return ExcelSheetData(
                source_path=path, source_sha256=hashlib.sha256(content).hexdigest(),
                source_size=len(content), sheet_name=sheet.title,
                header=header, rows=tuple(rows[1:]),
            )

    @classmethod
    def _validate_archive(cls, content: bytes) -> None:
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > cls.MAX_ARCHIVE_ENTRIES:
                    raise ValueError("Excelの内部ファイル数が読み込み上限を超えています。")
                if sum(entry.file_size for entry in entries) > cls.MAX_EXPANDED_SIZE:
                    raise ValueError("Excelの展開後サイズが読み込み上限を超えています。")
                if any(entry.flag_bits & 1 for entry in entries):
                    raise ValueError("暗号化されたExcelは読み込めません。")
        except zipfile.BadZipFile as error:
            raise ValueError("正常な.xlsxファイルとして読み込めません。") from error
