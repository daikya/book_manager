"""実際の.xlsxを使った読み取りの検証。"""
from datetime import datetime
import hashlib
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook

from book_manager.services.excel_workbook_reader_service import ExcelWorkbookReaderService


class TestExcelWorkbookReaderService(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "books.xlsx"
        self.reader = ExcelWorkbookReaderService()

    def write(self, rows, *, extra_sheet=False):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "読書記録"
        for row in rows:
            sheet.append(row)
        if extra_sheet:
            workbook.create_sheet("別シート")
        workbook.save(self.path)
        workbook.close()

    def cache_formula(self, coordinate, value, *, string=False):
        """Excelが保存する計算値をXMLへ入れ、数式を実行せず検証する。"""
        with zipfile.ZipFile(self.path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        name = "xl/worksheets/sheet1.xml"
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        root = ET.fromstring(entries[name])
        cell = next(cell for cell in root.iter(namespace + "c") if cell.get("r") == coordinate)
        if string:
            cell.set("t", "str")
        cached = cell.find(namespace + "v")
        if cached is None:
            cached = ET.SubElement(cell, namespace + "v")
        cached.text = str(value)
        entries[name] = ET.tostring(root)
        with zipfile.ZipFile(self.path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, content in entries.items():
                archive.writestr(name, content)

    def test_reads_headers_and_values_without_changing_file(self):
        self.write([["書籍名", "著者名"], ["本", "著者"]])
        original = self.path.read_bytes()
        result = self.reader.read_sheet(self.path)
        self.assertEqual(result.sheet_name, "読書記録")
        self.assertEqual(result.header.row_number, 1)
        self.assertEqual(result.rows[0].cells[0].coordinate, "A2")
        self.assertEqual(result.rows[0].cells[1].value, "著者")
        self.assertEqual(result.source_sha256, hashlib.sha256(original).hexdigest())
        self.assertEqual(result.source_size, len(original))
        self.assertEqual(original, self.path.read_bytes())

    def test_preserves_blank_rows_for_later_exclusion(self):
        self.write([["書籍名"], ["本"], [None], ["別の本"]])
        result = self.reader.read_sheet(self.path)
        self.assertEqual([row.row_number for row in result.rows], [2, 3, 4])
        self.assertTrue(result.rows[1].is_empty)
        self.assertFalse(result.rows[0].is_empty)

    def test_preserves_unused_and_duplicate_headers(self):
        self.write([["書籍名", None, "書籍名"], ["本", None, "別情報"]])
        result = self.reader.read_sheet(self.path)
        self.assertEqual([cell.value for cell in result.header.cells], ["書籍名", None, "書籍名"])

    def test_selects_header_row(self):
        self.write([["読書一覧"], [None], ["書籍名"], ["本"]])
        result = self.reader.read_sheet(self.path, header_row=3)
        self.assertEqual(result.header.row_number, 3)
        self.assertEqual(result.rows[0].row_number, 4)

    def test_rejects_invalid_header_row(self):
        self.write([["書籍名"]])
        for value in (0, -1, True, 1.5, "1"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.reader.read_sheet(self.path, header_row=value)

    def test_rejects_header_beyond_sheet(self):
        self.write([["書籍名"]])
        with self.assertRaises(ValueError):
            self.reader.read_sheet(self.path, header_row=3)

    def test_rejects_blank_header(self):
        self.write([[None], ["本"]])
        with self.assertRaises(ValueError):
            self.reader.read_sheet(self.path)

    def test_accepts_headers_without_data(self):
        self.write([["書籍名"]])
        self.assertEqual(self.reader.read_sheet(self.path).rows, ())

    def test_rejects_multiple_sheets(self):
        self.write([["書籍名"]], extra_sheet=True)
        with self.assertRaisesRegex(ValueError, "1枚"):
            self.reader.read_sheet(self.path)

    def test_keeps_missing_formula_value_as_cell_issue(self):
        self.write([["書籍名", "管理用計算"], ["本", "=1+1"]])
        cell = self.reader.read_sheet(self.path).rows[0].cells[1]
        self.assertTrue(cell.is_formula)
        self.assertIsNone(cell.value)
        self.assertIn("再計算", cell.issue)

    def test_formula_issue_does_not_become_empty_row(self):
        self.write([["書籍名"], ["=1+1"]])
        self.assertFalse(self.reader.read_sheet(self.path).rows[0].is_empty)

    def test_reads_numeric_cached_value_without_calculating_formula(self):
        self.write([["値"], ["=1+1"]])
        self.cache_formula("A2", 42)
        cell = self.reader.read_sheet(self.path).rows[0].cells[0]
        self.assertEqual(cell.value, 42)
        self.assertTrue(cell.is_formula)
        self.assertIsNone(cell.issue)

    def test_reads_string_cached_value(self):
        self.write([["書籍名"], ['="本"']])
        self.cache_formula("A2", "計算済み書籍名", string=True)
        cell = self.reader.read_sheet(self.path).rows[0].cells[0]
        self.assertEqual(cell.value, "計算済み書籍名")
        self.assertIsNone(cell.issue)

    def test_keeps_excel_error_as_cell_issue(self):
        self.write([["書籍名"], ["#DIV/0!"]])
        cell = self.reader.read_sheet(self.path).rows[0].cells[0]
        self.assertIn("エラー値", cell.issue)

    def test_preserves_date_and_display_format(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["読了日"])
        sheet["A2"] = datetime(2026, 5, 20)
        sheet["A2"].number_format = 'yyyy"年"m"月"'
        workbook.save(self.path)
        workbook.close()
        cell = self.reader.read_sheet(self.path).rows[0].cells[0]
        self.assertEqual(cell.value, datetime(2026, 5, 20))
        self.assertEqual(cell.number_format, 'yyyy"年"m"月"')

    def test_preserves_numbers_and_booleans_without_date_guessing(self):
        self.write([["数値", "真偽"], [45000, False]])
        row = self.reader.read_sheet(self.path).rows[0]
        self.assertEqual(row.cells[0].value, 45000)
        self.assertIs(row.cells[1].value, False)
        self.assertFalse(row.is_empty)

    def test_rejects_non_xlsx_extensions(self):
        for suffix in (".xls", ".xlsm", ".csv"):
            with self.subTest(suffix=suffix), self.assertRaises(ValueError):
                self.reader.read_sheet(self.path.with_suffix(suffix))

    def test_rejects_invalid_zip(self):
        self.path.write_bytes(b"not an xlsx")
        with self.assertRaises(ValueError):
            self.reader.read_sheet(self.path)

    def test_file_size_limit(self):
        self.write([["書籍名"]])
        with patch.object(self.reader, "MAX_FILE_SIZE", 10), self.assertRaises(ValueError):
            self.reader.read_sheet(self.path)

    def test_row_limit(self):
        self.write([["書籍名"], ["本"]])
        with patch.object(self.reader, "MAX_ROWS", 1), self.assertRaises(ValueError):
            self.reader.read_sheet(self.path)

    def test_column_limit(self):
        self.write([["書籍名", "著者名"]])
        with patch.object(self.reader, "MAX_COLUMNS", 1), self.assertRaises(ValueError):
            self.reader.read_sheet(self.path)

    def test_repeated_reads_release_file(self):
        self.write([["書籍名"], ["本"]])
        self.reader.read_sheet(self.path)
        self.reader.read_sheet(self.path)
        destination = self.path.with_name("renamed.xlsx")
        self.path.rename(destination)
        self.assertTrue(destination.is_file())
