"""列割り当てと全行検査の検証。DB登録やAPI呼び出しは許可しない。"""
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock

from openpyxl import Workbook

from book_manager.models.book_item import BookItem
from book_manager.models.excel_import import ExcelImportOptions
from book_manager.services.excel_import_preflight_service import ExcelImportPreflightService
from book_manager.services.excel_workbook_reader_service import ExcelWorkbookReaderService


class TestExcelImportPreflightService(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "books.xlsx"
        # prepare_book以外にアクセスすると失敗する。実際の入力検証は既存BookServiceの責務。
        self.books = Mock(spec=["prepare_book"])
        def prepare(**values):
            if not str(values.get("title", "")).strip():
                raise ValueError("書籍名は必須です。")
            values["title"] = values["title"].strip()
            return BookItem(**values)
        self.books.prepare_book.side_effect = prepare
        self.service = ExcelImportPreflightService(book_service=self.books)

    def write(self, headers, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(self.path)
        workbook.close()

    def preview(self, mapping, **options):
        return self.service.preview_file(
            self.path, options=ExcelImportOptions(tuple(mapping.items()), **options),
        )

    def test_title_and_author_without_isbn(self):
        self.write(["名前", "著者"], [[" 書籍 ", "著者名"]])
        result = self.preview({"title": 1, "author": 2})
        self.assertTrue(result.can_import)
        self.assertEqual(result.records[0].book.title, "書籍")
        self.assertIsNone(result.records[0].book.isbn)
        self.assertFalse(result.records[0].has_history)

    def test_months_create_history_and_suggest_completed_state(self):
        self.write(["本", "開始", "読了"], [["書籍", "平成21年9月", "平成21年10月"]])
        result = self.preview({"title": 1, "started_on": 2, "completed_on": 3})
        row = result.records[0]
        self.assertEqual((row.started_on, row.completed_on), ("2009-09", "2009-10"))
        self.assertEqual(row.book.completion_month, "2009-10")
        self.assertEqual(row.book.reading_status, "completed")
        self.assertTrue(row.reading_status_suggested)
        self.assertTrue(row.has_history)
        self.assertTrue(result.warnings)

    def test_started_only_suggests_reading_state(self):
        self.write(["本", "開始"], [["書籍", "2026-05"]])
        row = self.preview({"title": 1, "started_on": 2}).records[0]
        self.assertEqual(row.book.reading_status, "reading")
        self.assertIsNone(row.book.completion_month)

    def test_explicit_state_takes_priority_over_suggestion(self):
        self.write(["本", "読了", "状態"], [["書籍", "2026-05", "未読"]])
        row = self.preview({"title": 1, "completed_on": 2, "reading_status": 3}).records[0]
        self.assertEqual(row.book.reading_status, "unread")
        self.assertFalse(row.reading_status_suggested)

    def test_invalid_state_reports_cell(self):
        self.write(["本", "状態"], [["書籍", "不明"]])
        result = self.preview({"title": 1, "reading_status": 2})
        self.assertEqual(result.errors[0].coordinate, "B2")

    def test_same_title_author_and_isbn_are_not_merged(self):
        self.write(["本", "著者", "ISBN"], [["書籍", "著者", "9784101010014"]] * 2)
        result = self.preview({"title": 1, "author": 2, "isbn": 3})
        self.assertEqual(len(result.records), 2)
        self.assertTrue(result.can_import)

    def test_collects_errors_and_continues_to_last_row(self):
        self.write(["本", "読了"], [["正常", "2026-05"], ["不正", "2026-13"], ["最後", "2026-06"]])
        result = self.preview({"title": 1, "completed_on": 2})
        self.assertEqual(len(result.records), 2)
        self.assertEqual(result.errors[0].coordinate, "B3")
        self.assertFalse(result.can_import)

    def test_missing_title_is_error(self):
        self.write(["本", "著者"], [[None, "著者"]])
        result = self.preview({"title": 1, "author": 2})
        self.assertFalse(result.can_import)
        self.assertEqual(result.errors[0].row_number, 2)

    def test_unused_formula_is_not_an_import_error(self):
        self.write(["本", "集計"], [["書籍", "=1+1"]])
        result = self.preview({"title": 1})
        self.assertTrue(result.can_import)

    def test_mapped_formula_without_cache_is_error(self):
        self.write(["本", "著者"], [["書籍", '=CONCAT("著","者")']])
        result = self.preview({"title": 1, "author": 2})
        self.assertEqual(result.errors[0].coordinate, "B2")
        self.assertIn("再計算", result.errors[0].message)

    def test_mapped_excel_error_is_error(self):
        self.write(["本", "著者"], [["書籍", "#N/A"]])
        self.assertFalse(self.preview({"title": 1, "author": 2}).can_import)

    def test_explicitly_excluded_row_ignores_its_cell_error(self):
        self.write(["本"], [["#N/A"], ["正常"]])
        result = self.preview({"title": 1}, excluded_rows=(2,))
        self.assertTrue(result.can_import)
        self.assertEqual(result.skipped[0].row_number, 2)

    def test_marker_title_is_kept_unless_configured(self):
        self.write(["本"], [["【最終行】"]])
        self.assertEqual(len(self.preview({"title": 1}).records), 1)
        result = self.preview({"title": 1}, end_marker_title="【最終行】")
        self.assertFalse(result.can_import)
        self.assertEqual(len(result.skipped), 1)

    def test_marker_with_other_book_information_is_not_excluded(self):
        self.write(["本", "著者"], [["【最終行】", "著者"]])
        result = self.preview({"title": 1, "author": 2}, end_marker_title="【最終行】")
        self.assertTrue(result.can_import)

    def test_blank_assigned_cells_are_explicitly_skipped(self):
        self.write(["本", "管理番号"], [[None, 1], ["正常", 2]])
        result = self.preview({"title": 1})
        self.assertEqual(len(result.records), 1)
        self.assertEqual(result.skipped[0].row_number, 2)

    def test_reverse_dates_report_completed_cell(self):
        self.write(["本", "開始", "読了"], [["書籍", "2026-06", "2026-05"]])
        result = self.preview({"title": 1, "started_on": 2, "completed_on": 3})
        self.assertEqual(result.errors[0].coordinate, "C2")

    def test_mixed_precision_in_same_month_is_allowed(self):
        self.write(["本", "開始", "読了"], [["書籍", "2026-05-20", "2026-05"]])
        self.assertTrue(self.preview({"title": 1, "started_on": 2, "completed_on": 3}).can_import)

    def test_excel_date_with_month_only_format_keeps_month_precision(self):
        self.write(["本", "読了"], [["書籍", datetime(2026, 5, 20)]])
        from openpyxl import load_workbook
        workbook = load_workbook(self.path)
        workbook.active["B2"].number_format = 'yyyy"年"m"月"'
        workbook.save(self.path)
        workbook.close()
        row = self.preview({"title": 1, "completed_on": 2}).records[0]
        self.assertEqual(row.completed_on, "2026-05")

    def test_excel_full_date_keeps_day_precision(self):
        self.write(["本", "読了"], [["書籍", datetime(2026, 5, 20)]])
        self.assertEqual(self.preview({"title": 1, "completed_on": 2}).records[0].completed_on, "2026-05-20")

    def test_numeric_date_is_not_guessed(self):
        self.write(["本", "読了"], [["書籍", 45000]])
        self.assertFalse(self.preview({"title": 1, "completed_on": 2}).can_import)

    def test_conflicting_completion_month_is_rejected(self):
        self.write(["本", "読了", "読了年月"], [["書籍", "2026-05", "2026-06"]])
        result = self.preview({"title": 1, "completed_on": 2, "completion_month": 3})
        self.assertEqual(result.errors[0].coordinate, "C2")

    def test_text_multiline_tags_favorite_and_rating(self):
        self.write(["本", "紹介", "メモ", "タグ", "お気に入り", "評価"],
                   [["書籍", "一行目\n二行目", "感想\n続き", "音声|既読", True, 5]])
        row = self.preview({"title": 1, "description": 2, "history_notes": 3,
                            "tags": 4, "is_favorite": 5, "rating": 6}).records[0]
        self.assertEqual(row.book.description, "一行目\n二行目")
        self.assertEqual(row.history_notes, "感想\n続き")
        self.assertTrue(row.has_history)
        self.assertTrue(row.book.is_favorite)
        self.assertEqual(row.book.rating, 5)
        self.assertEqual(row.book.tags, ("音声", "既読"))

    def test_invalid_rating_and_favorite_collect_both_cell_errors(self):
        self.write(["本", "評価", "お気に入り"], [["書籍", 6, "不明"]])
        result = self.preview({"title": 1, "rating": 2, "is_favorite": 3})
        self.assertEqual({issue.coordinate for issue in result.errors}, {"B2", "C2"})

    def test_default_publication_format_only_when_column_not_assigned(self):
        self.write(["本", "出版形態"], [["書籍", None]])
        row = self.preview({"title": 1}, default_publication_format="音声版").records[0]
        self.assertEqual(row.book.publication_format, "音声版")
        row = self.preview({"title": 1, "publication_format": 2}, default_publication_format="音声版").records[0]
        self.assertIsNone(row.book.publication_format)

    def test_invalid_column_mapping_is_rejected(self):
        for mapping in ((('author', 1),), (('title', 0),), (('title', 1), ('author', 1)),
                        (('title', 1), ('title', 2)), (('title', 1), ('unknown', 2))):
            with self.subTest(mapping=mapping), self.assertRaises(ValueError):
                ExcelImportOptions(mapping)

    def test_missing_column_or_excluded_row_is_rejected(self):
        self.write(["本"], [["書籍"]])
        with self.assertRaises(ValueError):
            self.preview({"title": 2})
        with self.assertRaises(ValueError):
            self.preview({"title": 1}, excluded_rows=(100,))

    def test_source_is_unchanged_and_reader_snapshot_is_retained(self):
        self.write(["本"], [["書籍"]])
        original = self.path.read_bytes()
        sheet = ExcelWorkbookReaderService().read_sheet(self.path)
        result = self.service.preview_sheet(sheet, options=ExcelImportOptions((("title", 1),)))
        self.assertEqual(result.source_sha256, sheet.source_sha256)
        self.assertEqual(self.path.read_bytes(), original)
        self.books.prepare_book.assert_called_once()
